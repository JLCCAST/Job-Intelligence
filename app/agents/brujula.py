"""
brujula.py -- agente Brújula

Tarea fija: decirte qué estudiar primero, a partir de TODAS tus
ofertas (las que trajo Sabueso y las que analizaste tú).

La prioridad se CALCULA en Python, no la decide Gemini:

    prioridad = demanda × distancia × tendencia

    demanda   = qué tan seguido la piden: Σ peso / n° de ofertas
                (peso 1.0 si es obligatoria, 0.5 si es deseable)
    distancia = qué tan lejos estás:
                1.0  no la tienes
                0.5  la tienes en un nivel menor al más alto que piden
                0.25 está en tu CV pero no en tu perfil (user_skill)
                0    la cubres -> no se recomienda
    tendencia = 1 + cambio de demanda entre los últimos 30 días y
                los 30 anteriores (entre 0.5 y 1.5). Si alguna de
                las dos ventanas tiene pocas ofertas, vale 1.0 y se
                marca "sin datos suficientes".

Gemini solo entra al final, para convertir las 5 primeras en un plan
de estudio con un proyecto pequeño por tecnología.
"""

import json
from dataclasses import asdict, dataclass
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Optional

from pydantic import BaseModel, Field

from app.cv.cargador import RAIZ
from app.gemini import generar_estructurado
from app.scoring import ORDEN_NIVEL

PESO_TIPO = {"required": 1.0, "preferred": 0.5}
VENTANA_DIAS = 30
MIN_OFERTAS_POR_VENTANA = 5
MUESTRA_MINIMA = 10
TOP_PLAN = 5
# Las habilidades blandas no se "estudian" con un curso: fuera del ranking.
CATEGORIAS_EXCLUIDAS = {"Soft_Skill"}

DISTANCIA_NO_LA_TIENES = 1.0
DISTANCIA_NIVEL_MENOR = 0.5
DISTANCIA_SOLO_EN_CV = 0.25

RUTA_PLAN = RAIZ / "data" / "brujula_plan.json"
CARPETA_INFORMES = RAIZ / "salidas" / "brujula"


@dataclass
class Prioridad:
    tecnologia: str
    categoria: str
    ofertas: int            # en cuántas ofertas aparece
    obligatoria_en: int
    demanda: float          # 0..1
    distancia: float
    estado: str             # "Te falta" / "Nivel menor (piden X)" / "En tu CV, no en tu perfil"
    tendencia: float
    tendencia_txt: str      # "sube" / "baja" / "estable" / "nueva" / "sin datos suficientes"
    prioridad: float


# --------------------------------------------------------------
# Cálculo (sin LLM)
# --------------------------------------------------------------

def _nivel(n: Optional[str]) -> int:
    return ORDEN_NIVEL.get(n or "", 0)


def _a_fecha(valor) -> date:
    if isinstance(valor, datetime):
        return valor.astimezone(timezone.utc).date() if valor.tzinfo else valor.date()
    if isinstance(valor, date):
        return valor
    return datetime.fromisoformat(str(valor)).date()


def _tendencia(filas_tech: list[dict], ofertas_por_ventana: tuple[set, set]) -> tuple[float, str]:
    recientes, anteriores = ofertas_por_ventana
    if len(recientes) < MIN_OFERTAS_POR_VENTANA or len(anteriores) < MIN_OFERTAS_POR_VENTANA:
        return 1.0, "sin datos suficientes"

    jobs = {f["job_id"] for f in filas_tech}
    share_reciente = len(jobs & recientes) / len(recientes)
    share_anterior = len(jobs & anteriores) / len(anteriores)
    if share_anterior == 0 and share_reciente > 0:
        return 1.5, "nueva"

    cambio = max(-0.5, min(0.5, share_reciente - share_anterior))
    texto = "sube" if cambio > 0.05 else "baja" if cambio < -0.05 else "estable"
    return round(1 + cambio, 3), texto


def calcular_prioridades(
    filas: list[dict],
    skills_cv: set[str] = frozenset(),
    respuestas_confirmadas: Optional[dict[str, Optional[str]]] = None,
    hoy: Optional[date] = None,
) -> list[Prioridad]:
    """
    `filas`: consultas.requisitos_para_brujula().
    `skills_cv`: CVBase.skills_declaradas() (minúsculas).
    `respuestas_confirmadas`: {tecnología en minúsculas: nivel} de lo
        que confirmaste tener en el Preguntón.
    """
    respuestas_confirmadas = respuestas_confirmadas or {}
    hoy = hoy or date.today()
    total_ofertas = len({f["job_id"] for f in filas})
    if total_ofertas == 0:
        return []

    corte_reciente = hoy - timedelta(days=VENTANA_DIAS)
    corte_anterior = hoy - timedelta(days=2 * VENTANA_DIAS)
    recientes, anteriores = set(), set()
    for f in filas:
        fecha = _a_fecha(f["fecha_registro"])
        if fecha > corte_reciente:
            recientes.add(f["job_id"])
        elif fecha > corte_anterior:
            anteriores.add(f["job_id"])

    por_tech: dict[str, list[dict]] = {}
    for f in filas:
        if f["categoria"] not in CATEGORIAS_EXCLUIDAS:
            por_tech.setdefault(f["tecnologia"], []).append(f)

    resultado = []
    for tech, filas_tech in por_tech.items():
        clave = tech.strip().lower()
        mi_nivel = filas_tech[0]["mi_nivel"] or respuestas_confirmadas.get(clave)
        tengo = filas_tech[0]["mi_nivel"] is not None or clave in respuestas_confirmadas
        pedido_max = max(filas_tech, key=lambda f: _nivel(f["nivel_pedido"]))["nivel_pedido"]

        if tengo:
            if mi_nivel and _nivel(pedido_max) > _nivel(mi_nivel):
                distancia, estado = DISTANCIA_NIVEL_MENOR, f"Nivel menor (piden {pedido_max})"
            elif not mi_nivel and _nivel(pedido_max) > 0:
                distancia, estado = DISTANCIA_NIVEL_MENOR, f"Sin nivel declarado (piden {pedido_max})"
            else:
                continue  # la cubres
        elif clave in skills_cv:
            distancia, estado = DISTANCIA_SOLO_EN_CV, "En tu CV, no en tu perfil"
        else:
            distancia, estado = DISTANCIA_NO_LA_TIENES, "Te falta"

        # Una oferta puede listar la misma tecnología una sola vez
        # (PK job_id+technology_id), así que sumar por fila es sumar por oferta.
        demanda = sum(PESO_TIPO.get(f["requirement_type"], 0.5) for f in filas_tech) / total_ofertas
        tendencia, tendencia_txt = _tendencia(filas_tech, (recientes, anteriores))

        resultado.append(Prioridad(
            tecnologia=tech,
            categoria=filas_tech[0]["categoria"],
            ofertas=len(filas_tech),
            obligatoria_en=sum(1 for f in filas_tech if f["requirement_type"] == "required"),
            demanda=round(demanda, 3),
            distancia=distancia,
            estado=estado,
            tendencia=tendencia,
            tendencia_txt=tendencia_txt,
            prioridad=round(demanda * distancia * tendencia, 4),
        ))

    resultado.sort(key=lambda p: (-p.prioridad, -p.ofertas, p.tecnologia.lower()))
    return resultado


# --------------------------------------------------------------
# Plan de estudio (Gemini)
# --------------------------------------------------------------

class ItemPlan(BaseModel):
    tecnologia: str = Field(description="Exactamente uno de los nombres de la lista entregada")
    por_que: str = Field(description="Una frase: por qué conviene, citando los datos entregados")
    que_aprender: list[str] = Field(description="3 a 5 temas concretos, de lo básico a lo aplicado")
    proyecto: str = Field(description="Un proyecto pequeño para el portafolio que lo demuestre")
    semanas: int = Field(description="Semanas estimadas a ritmo de estudiante (5-8 h por semana)")


class PlanEstudio(BaseModel):
    resumen: str = Field(description="2 o 3 frases con el orden recomendado y por qué")
    items: list[ItemPlan]


PROMPT_PLAN = """\
Eres un mentor de carrera en datos para un estudiante de Ingeniería de
Sistemas en Perú. Arma un plan de estudio para estas tecnologías, en
este orden de prioridad (ya calculado a partir de ofertas reales):

{prioridades}

Lo que el estudiante YA sabe (úsalo como punto de partida):
{base}

Reglas:
- Un ítem por cada tecnología de la lista, con el mismo nombre exacto.
- "por_que" debe citar los datos entregados (ofertas, obligatoria, tendencia).
  No inventes estadísticas del mercado.
- No inventes links ni nombres de cursos específicos. Si recomiendas
  recursos, di el tipo ("documentación oficial", "tutorial oficial").
- El proyecto debe poder hacerse en local y aprovechar lo que ya sabe.
- Escribe en español, directo y sin adornos.
"""


def armar_prompt_plan(top: list[Prioridad], skills_base: list[str]) -> str:
    lineas = [
        f"{i}. {p.tecnologia} ({p.categoria}): en {p.ofertas} oferta(s), obligatoria en "
        f"{p.obligatoria_en}; estado: {p.estado}; tendencia: {p.tendencia_txt}"
        for i, p in enumerate(top, 1)
    ]
    return PROMPT_PLAN.format(prioridades="\n".join(lineas), base=", ".join(skills_base) or "(sin datos)")


def generar_plan(top: list[Prioridad], skills_base: list[str]) -> tuple[Optional[PlanEstudio], Optional[str]]:
    if not top:
        return None, "No hay nada que estudiar: cubres todo lo que piden tus ofertas."
    resultado = generar_estructurado(armar_prompt_plan(top, skills_base), PlanEstudio, agente="brujula")
    if not resultado.success:
        return None, resultado.error

    # Solo se aceptan ítems de la lista calculada: Gemini no elige qué estudiar.
    permitidas = {p.tecnologia.lower() for p in top}
    plan = resultado.data
    plan.items = [i for i in plan.items if i.tecnologia.strip().lower() in permitidas]
    return plan, None


def guardar_plan(plan: PlanEstudio, top: list[Prioridad], ruta: Optional[Path] = None) -> None:
    ruta = ruta or RUTA_PLAN
    ruta.parent.mkdir(parents=True, exist_ok=True)
    datos = {
        "generado": datetime.now().isoformat(timespec="minutes"),
        "prioridades": [asdict(p) for p in top],
        "plan": plan.model_dump(),
    }
    ruta.write_text(json.dumps(datos, ensure_ascii=False, indent=1), encoding="utf-8")


def cargar_plan(ruta: Optional[Path] = None) -> Optional[dict]:
    ruta = ruta or RUTA_PLAN
    if not ruta.exists():
        return None
    return json.loads(ruta.read_text(encoding="utf-8"))


# --------------------------------------------------------------
# Informe DOCX
# --------------------------------------------------------------

def generar_informe(
    prioridades: list[Prioridad],
    total_ofertas: int,
    plan: Optional[dict] = None,
    carpeta: Optional[Path] = None,
) -> Path:
    """Informe con el ranking y, si existe, el plan de estudio."""
    # Se reutiliza el formato de Imprenta para que todos los DOCX se vean igual.
    from docx.shared import Pt

    from app.agents.imprenta import COLOR_TITULOS, _documento_base, _etiqueta_valor, _titulo_seccion, _vineta

    doc = _documento_base()
    p = doc.add_paragraph()
    run = p.add_run(f"Brújula — qué estudiar ({date.today().isoformat()})")
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = COLOR_TITULOS
    doc.add_paragraph(f"Calculado con {total_ofertas} oferta(s) registrada(s), sin contar las descartadas.")
    if total_ofertas < MUESTRA_MINIMA:
        doc.add_paragraph(
            f"Con menos de {MUESTRA_MINIMA} ofertas el ranking es orientativo: refleja tus ofertas, "
            "no el mercado."
        )

    _titulo_seccion(doc, "Prioridades")
    if not prioridades:
        doc.add_paragraph("Cubres todo lo que piden tus ofertas.")
    tabla = doc.add_table(rows=1, cols=5)
    tabla.style = "Light Grid Accent 1"
    for celda, texto in zip(tabla.rows[0].cells, ("Tecnología", "Ofertas", "Estado", "Tendencia", "Prioridad")):
        celda.text = texto
    for pr in prioridades[:15]:
        c = tabla.add_row().cells
        c[0].text = pr.tecnologia
        c[1].text = f"{pr.ofertas} ({pr.obligatoria_en} oblig.)"
        c[2].text = pr.estado
        c[3].text = pr.tendencia_txt
        c[4].text = f"{pr.prioridad:.2f}"

    if plan:
        _titulo_seccion(doc, "Plan de estudio")
        doc.add_paragraph(plan["plan"]["resumen"])
        for item in plan["plan"]["items"]:
            _etiqueta_valor(doc, f"{item['tecnologia']} ", f"(~{item['semanas']} semanas)")
            doc.add_paragraph(item["por_que"])
            for tema in item["que_aprender"]:
                _vineta(doc, tema)
            _etiqueta_valor(doc, "Proyecto: ", item["proyecto"])

    carpeta = carpeta or CARPETA_INFORMES
    carpeta.mkdir(parents=True, exist_ok=True)
    ruta = carpeta / f"Brujula_{date.today().isoformat()}.docx"
    doc.save(ruta)
    return ruta
