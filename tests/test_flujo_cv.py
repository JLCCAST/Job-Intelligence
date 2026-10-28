"""
Tests del flujo de adaptar CV: Preguntón -> Pluma -> Imprenta.

Se usa el CV de ejemplo (data/cv_base.example.yaml), una oferta
falsa y un Gemini falso; el checkpointer es en memoria y los
archivos se escriben en una carpeta temporal.
"""

import pytest
from docx import Document
from langgraph.checkpoint.memory import InMemorySaver

from app.agents import balanza, batuta, imprenta, pluma, pregunton
from app.cv import respuestas as modulo_respuestas
from app.cv.cargador import RAIZ, cargar_cv_base
from app.cv.modelo import CVBase
from app.gemini import ResultadoEstructurado

CV_EJEMPLO = cargar_cv_base(RAIZ / "data" / "cv_base.example.yaml")

OFERTA = {
    "job_id": 42,
    "puesto": "Analista de Datos Jr",
    "empresa": "Empresa Ficticia",
    "area": "Data_Analytics",
    "match_score": 55.0,
    "cumple_ingles": True,
    "requisitos": [
        {"nombre": "Python", "requirement_type": "required", "nivel_pedido": "intermedio",
         "nivel_usuario": "intermedio", "bajo_nivel": False},
        {"nombre": "Power BI", "requirement_type": "required", "nivel_pedido": "unspecified",
         "nivel_usuario": None, "bajo_nivel": False},
        {"nombre": "Airflow", "requirement_type": "required", "nivel_pedido": "basico",
         "nivel_usuario": None, "bajo_nivel": False},
        {"nombre": "dbt", "requirement_type": "preferred", "nivel_pedido": "unspecified",
         "nivel_usuario": None, "bajo_nivel": False},
        {"nombre": "Inglés", "requirement_type": "required", "nivel_pedido": "intermedio",
         "nivel_usuario": None, "bajo_nivel": False},
    ],
}

REDACCION_LIMPIA = {
    "perfil": "Estudiante de Ingeniería de Sistemas orientado a BI con SQL, Python, Power BI y dbt.",
    "experiencia": [{"indice": 0, "logros": ["Automaticé en Python un reporte semanal que tomaba 3 horas manuales."]}],
    "proyectos": [{"indice": 0, "logros": ["Modelé un esquema estrella en SQL Server y un dashboard en Power BI con 8 KPIs."]}],
    "habilidades_destacadas": ["Power BI", "SQL", "dbt"],
}
REDACCION_CON_AIRFLOW = {**REDACCION_LIMPIA, "perfil": "Analista con experiencia en Airflow y Python."}


# --------- Preguntón ---------

def test_pregunton_solo_pregunta_lo_que_no_esta_en_perfil_cv_ni_respuestas():
    preguntas = pregunton.armar_preguntas(
        OFERTA["requisitos"], skills_cv=CV_EJEMPLO.skills_declaradas(), ya_respondidas=set()
    )
    # Python: en el perfil. Power BI: en el CV. Inglés: no se pregunta.
    assert [p["tecnologia"] for p in preguntas] == ["Airflow", "dbt"]
    assert "obligatorio (nivel basico)" in preguntas[0]["texto"]


def test_pregunton_no_repite_lo_ya_respondido():
    preguntas = pregunton.armar_preguntas(
        OFERTA["requisitos"], skills_cv=CV_EJEMPLO.skills_declaradas(), ya_respondidas={"airflow"}
    )
    assert [p["tecnologia"] for p in preguntas] == ["dbt"]


# --------- Pluma ---------

def test_seleccion_pone_primero_el_proyecto_mas_relevante():
    cv = CVBase.model_validate({
        "contacto": {"nombre": "X", "email": "x@x.com"},
        "perfil": "p",
        "proyectos": [
            {"nombre": "Planilla", "logros": ["Hice una planilla."], "stack": ["Excel"]},
            {"nombre": "Dashboard", "logros": ["Otro logro.", "Dashboard en Power BI."], "stack": ["Power BI"]},
        ],
    })
    seleccion = pluma.calcular_seleccion(cv, OFERTA)
    assert [p["indice"] for p in seleccion["proyectos"]] == [1, 0]
    # Dentro del proyecto, el bullet que menciona la oferta va primero.
    assert seleccion["proyectos"][0]["logros"] == [1, 0]


def _seleccion_ejemplo():
    return pluma.calcular_seleccion(CV_EJEMPLO, OFERTA)


def test_validador_acepta_redaccion_limpia():
    confirmadas = [{"tecnologia": "dbt", "la_tengo": True}]
    assert pluma.encontrar_violaciones(CV_EJEMPLO, OFERTA, _seleccion_ejemplo(), REDACCION_LIMPIA, confirmadas) == []


def test_validador_rechaza_tecnologia_no_confirmada():
    violaciones = pluma.encontrar_violaciones(CV_EJEMPLO, OFERTA, _seleccion_ejemplo(), REDACCION_CON_AIRFLOW, [])
    assert any("Airflow" in v for v in violaciones)


def test_validador_rechaza_cifras_inventadas():
    inventada = {**REDACCION_LIMPIA, "proyectos": [{"indice": 0, "logros": ["Reduje el tiempo de reporte en 45%."]}]}
    confirmadas = [{"tecnologia": "dbt", "la_tengo": True}]
    violaciones = pluma.encontrar_violaciones(CV_EJEMPLO, OFERTA, _seleccion_ejemplo(), inventada, confirmadas)
    assert any("45" in v for v in violaciones)


def test_validador_rechaza_indices_cambiados():
    cambiada = {**REDACCION_LIMPIA, "proyectos": [{"indice": 3, "logros": ["x"]}]}
    confirmadas = [{"tecnologia": "dbt", "la_tengo": True}]
    violaciones = pluma.encontrar_violaciones(CV_EJEMPLO, OFERTA, _seleccion_ejemplo(), cambiada, confirmadas)
    assert any("índices" in v for v in violaciones)


# --------- Flujo completo con Batuta ---------

@pytest.fixture
def entorno(monkeypatch, tmp_path):
    """CV de ejemplo, oferta falsa, Gemini falso, archivos en tmp."""
    redacciones = []  # cola de lo que "responde" Gemini, en orden

    def fake_gemini(prompt, schema, agente):
        datos = redacciones.pop(0) if len(redacciones) > 1 else redacciones[0]
        return ResultadoEstructurado(success=True, data=schema.model_validate(datos))

    for modulo in (pregunton, pluma, imprenta):
        monkeypatch.setattr(modulo, "cargar_cv_base", lambda: CV_EJEMPLO)
    monkeypatch.setattr(balanza, "leer_oferta_para_cv", lambda job_id: OFERTA)
    monkeypatch.setattr(pluma, "generar_estructurado", fake_gemini)
    monkeypatch.setattr(modulo_respuestas, "RUTA_RESPUESTAS", tmp_path / "respuestas.yaml")
    monkeypatch.setattr(imprenta, "CARPETA_SALIDAS", tmp_path / "salidas")

    grafo = batuta.construir_grafo_cv(InMemorySaver())
    return grafo, redacciones, tmp_path


def _texto_docx(ruta) -> str:
    doc = Document(ruta)
    partes = [p.text for p in doc.paragraphs]
    partes += [c.text for t in doc.tables for fila in t.rows for c in fila.cells]
    return "\n".join(partes)


RESPUESTAS = [
    {"tecnologia": "Airflow", "la_tengo": False},
    {"tecnologia": "dbt", "la_tengo": True, "nivel": "basico", "donde": "Curso de dbt"},
]


def test_flujo_pausa_corrige_y_genera_los_dos_docx(entorno):
    grafo, redacciones, tmp_path = entorno
    # 1er intento inventa Airflow -> el validador lo rechaza -> 2do intento limpio.
    redacciones += [REDACCION_CON_AIRFLOW, REDACCION_LIMPIA]

    r = batuta.iniciar_cv(42, grafo=grafo)
    assert r["esperando_respuestas"]
    assert [p["tecnologia"] for p in r["preguntas"]] == ["Airflow", "dbt"]

    r = batuta.responder_cv(r["thread_id"], RESPUESTAS, grafo=grafo)
    estado = r["estado"]
    assert not r["esperando_respuestas"]
    assert estado["intentos"] == 2
    assert estado["violaciones"] == []

    cv = _texto_docx(estado["ruta_cv"])
    assert "Modelé un esquema estrella" in cv          # texto de Pluma
    assert "Airflow" not in cv
    assert estado["ruta_cv"].endswith("CV_Empresa_Ficticia_Analista_de_Datos_Jr.docx")

    brechas = _texto_docx(estado["ruta_brechas"])
    assert "Airflow (obligatorio): no aparece en tu perfil." in brechas
    assert "Confirmado en Preguntón" in brechas        # dbt
    assert "En tu CV (no en tu perfil)" in brechas     # Power BI

    # Las respuestas quedaron guardadas: la siguiente vez no pregunta.
    assert len(modulo_respuestas.cargar_respuestas()) == 2
    r2 = batuta.iniciar_cv(42, grafo=grafo)
    assert not r2["esperando_respuestas"]


def test_si_pluma_nunca_pasa_el_validador_usa_tus_bullets_originales(entorno):
    grafo, redacciones, _ = entorno
    redacciones.append(REDACCION_CON_AIRFLOW)  # siempre inventa

    r = batuta.iniciar_cv(42, grafo=grafo)
    r = batuta.responder_cv(r["thread_id"], RESPUESTAS, grafo=grafo)
    estado = r["estado"]

    assert estado["intentos"] == pluma.MAX_INTENTOS
    assert estado["violaciones"]
    cv = _texto_docx(estado["ruta_cv"])
    assert "Airflow" not in cv
    assert CV_EJEMPLO.proyectos[0].logros[0] in cv     # bullet original
    assert "no pasó el validador" in _texto_docx(estado["ruta_brechas"])
